from __future__ import annotations
import fnmatch
import json
import ntpath
import os
import posixpath
import re
import shlex
import stat
import subprocess
import sys
import time
import unicodedata
from concurrent.futures import ThreadPoolExecutor
from enum import Enum
from pathlib import Path
from typing import Any, Callable, Iterable
GOOGLE_WORKSPACE_EXTENSIONS: frozenset[str] = frozenset()


def workspace(*args, **kwargs):
    return None


def enabled() -> bool:
    return False

from graphify.paths import GRAPHIFY_OUT, GRAPHIFY_OUT_NAME, out


class FileType(str, Enum):
    CODE = 'code'
    DOCUMENT = 'document'
    PAPER = 'paper'
    IMAGE = 'image'
    VIDEO = 'video'
_MANIFEST_PATH = str(out('manifest.json'))
_MTIME_COARSE_S = 2.0
_MTIME_SUBSECOND_S = 0.05
CODE_EXTENSIONS = {'.py', '.ts', '.tsx', '.mts', '.cts', '.js', '.jsx', '.mjs', '.cjs', '.ejs', '.ets', '.go', '.rs', '.vb', '.cbl', '.cob', '.cobol', '.cpy', '.java', '.groovy', '.gradle', '.cpp', '.cc', '.cxx', '.c', '.h', '.hpp', '.cu', '.cuh', '.metal', '.rb', '.rake', '.swift', '.kt', '.kts', '.cs', '.scala', '.php', '.lua', '.luau', '.toc', '.zig', '.ps1', '.psm1', '.psd1', '.ex', '.exs', '.m', '.mm', '.ml', '.mli', '.jl', '.vue', '.svelte', '.astro', '.dart', '.v', '.sv', '.svh', '.vh', '.sql', '.r', '.f', '.F', '.f90', '.F90', '.f95', '.F95', '.f03', '.F03', '.f08', '.F08', '.pas', '.pp', '.dpr', '.dpk', '.lpr', '.inc', '.dfm', '.lfm', '.lpk', '.sh', '.bash', '.json', '.tf', '.tfvars', '.hcl', '.dm', '.dme', '.dmi', '.dmm', '.dmf', '.sln', '.slnx', '.csproj', '.fsproj', '.vbproj', '.xaml', '.razor', '.cshtml', '.cls', '.trigger', '.lisp', '.cl', '.lsp', '.asd', '.robot', '.resource', '.sol', '.erl', '.hrl', '.escript'}
DOC_EXTENSIONS = {'.md', '.mdx', '.qmd', '.skill', '.txt', '.rst', '.html', '.yaml', '.yml'}
PAPER_EXTENSIONS = {'.pdf'}
IMAGE_EXTENSIONS = {'.png', '.jpg', '.jpeg', '.gif', '.webp', '.svg'}
OFFICE_EXTENSIONS = {'.docx', '.xlsx'}
VIDEO_EXTENSIONS = {'.mp4', '.mov', '.webm', '.mkv', '.avi', '.m4v', '.mp3', '.wav', '.m4a', '.ogg'}
CORPUS_WARN_THRESHOLD = 50000
CORPUS_UPPER_THRESHOLD = 500000
FILE_COUNT_UPPER = 500
_OFFICE_MAX_RAW_BYTES = 50 * 1024 * 1024
_OFFICE_MAX_DECOMPRESSED_BYTES = 512 * 1024 * 1024
_OFFICE_MAX_COMPRESSION_RATIO = 200


def cap(path: Path, cap: int=_OFFICE_MAX_RAW_BYTES) -> bool:
    try:
        return path.stat().st_size <= cap
    except OSError:
        return False


def zipcap(path: Path) -> bool:
    import zipfile

    if not cap(path):
        return False

    try:
        with zipfile.ZipFile(path) as zf:
            infos = zf.infolist()
            compressed = sum((i.compress_size for i in infos)) or 1
            declared = sum((i.file_size for i in infos))

            if declared > _OFFICE_MAX_DECOMPRESSED_BYTES:
                return False

            if declared / compressed > _OFFICE_MAX_COMPRESSION_RATIO:
                return False

            total = 0

            for info in infos:
                with zf.open(info) as member:
                    while True:
                        chunk = member.read(1024 * 1024)

                        if not chunk:
                            break

                        total += len(chunk)

                        if total > _OFFICE_MAX_DECOMPRESSED_BYTES:
                            return False
    except (zipfile.BadZipFile, OSError, EOFError):
        return False

    return True

_CREDENTIAL_STORE_DIRS = frozenset({'.ssh', '.gnupg', '.aws', '.gcloud'})
_AMBIGUOUS_SENSITIVE_DIRS = frozenset({'secrets', '.secrets', 'credentials'})
_SENSITIVE_PATTERNS = [re.compile('(^|[\\\\/])\\.(env|envrc)(\\.|$)', re.IGNORECASE), re.compile('\\.(pem|key|p12|pfx|cert|crt|der|p8)$', re.IGNORECASE), re.compile('(^|[^A-Za-z0-9])(id_rsa|id_dsa|id_ecdsa|id_ed25519)(\\.pub)?$', re.IGNORECASE), re.compile('^secring(\\.(gpg|pgp))?$', re.IGNORECASE), re.compile('(\\.netrc|\\.pgpass|\\.htpasswd|\\.npmrc|\\.pypirc|\\.git-credentials|\\.boto)$', re.IGNORECASE)]
_ENV_TEMPLATE_SUFFIXES = ('.example', '.sample', '.template', '.dist')


def template(name: str) -> bool:
    lower = name.lower()

    if not lower.endswith(_ENV_TEMPLATE_SUFFIXES):
        return False

    return bool(re.match('\\.(env|envrc)\\.', lower))

_GENERIC_KEYWORD_PATTERNS = [re.compile('(?<![a-zA-Z0-9])(credential|secret|passwd|password|private_key)s?(?![a-zA-Z])', re.IGNORECASE), re.compile('(?<![a-zA-Z0-9])tokens?(?![a-zA-Z])', re.IGNORECASE), re.compile('(?<![a-zA-Z0-9])service[._-]?account(?![a-zA-Z])', re.IGNORECASE)]
_PROSE_EXTS = frozenset({'.md', '.markdown', '.rst', '.org', '.adoc', '.tex'})
_BARE_PLURAL_TOKENS = re.compile('(?<![a-zA-Z0-9])tokens(?![a-zA-Z])', re.IGNORECASE)
_SECRET_PRONE_DATA_EXTS = frozenset({'.json', '.yaml', '.yml', '.toml', '.ini', '.cfg', '.conf', '.config', '.xml', '.properties', '.env', '.txt', '.tfvars'})
_WORD_SPLIT = re.compile('[-_\\s.]+')


def prose(path: Path) -> bool:
    if path.suffix.lower() not in _PROSE_EXTS:
        return False

    stem = Path(path.name).stem.lstrip('.') or Path(path.name).stem

    if _BARE_PLURAL_TOKENS.fullmatch(stem):
        return True

    return not any((p.fullmatch(stem) for p in _GENERIC_KEYWORD_PATTERNS))


def keyword(name: str) -> bool:
    stem = Path(name).stem.lstrip('.') or Path(name).stem

    for pat in _GENERIC_KEYWORD_PATTERNS:
        hit = False

        for m in pat.finditer(stem):
            hit = True

            if m.end() == len(stem):
                return True

        if hit and len([w for w in _WORD_SPLIT.split(stem) if w]) <= 2:
            return True

    return False

_PAPER_SIGNALS = [re.compile('\\barxiv\\b', re.IGNORECASE), re.compile('\\bdoi\\s*:', re.IGNORECASE), re.compile('\\babstract\\b', re.IGNORECASE), re.compile('\\bproceedings\\b', re.IGNORECASE), re.compile('\\bjournal\\b', re.IGNORECASE), re.compile('\\bpreprint\\b', re.IGNORECASE), re.compile('\\\\cite\\{'), re.compile('\\[\\d+\\]'), re.compile('\\[\\n\\d+\\n\\]'), re.compile('eq\\.\\s*\\d+|equation\\s+\\d+', re.IGNORECASE), re.compile('\\d{4}\\.\\d{4,5}'), re.compile('\\bwe propose\\b', re.IGNORECASE), re.compile('\\bliterature\\b', re.IGNORECASE)]
_PAPER_SIGNAL_THRESHOLD = 3


def graphable(path: Path) -> bool:
    return classify(path) == FileType.CODE and path.suffix.lower() not in _SECRET_PRONE_DATA_EXTS


def sensitive(path: Path) -> bool:
    parents = path.parts[:-1]

    if any((part.lower() in _CREDENTIAL_STORE_DIRS for part in parents)):
        return True

    if any((part.lower() in _AMBIGUOUS_SENSITIVE_DIRS for part in parents)) and (not graphable(path)):
        return True

    name = path.name

    if any((p.search(name) for p in _SENSITIVE_PATTERNS)) and (not template(name)):
        return True

    if keyword(name):
        return not (graphable(path) or prose(path))

    return False


def paper(path: Path) -> bool:
    try:
        text = path.read_text(encoding='utf-8', errors='ignore')[:3000]
        hits = sum((1 for pattern in _PAPER_SIGNALS if pattern.search(text)))

        return hits >= _PAPER_SIGNAL_THRESHOLD
    except Exception:
        return False

_ASSET_DIR_MARKERS = {'.imageset', '.xcassets', '.appiconset', '.colorset', '.launchimage'}
_SHEBANG_CODE_INTERPRETERS = {'python', 'python3', 'python2', 'ruby', 'perl', 'node', 'nodejs', 'bash', 'sh', 'dash', 'zsh', 'fish', 'ksh', 'tcsh', 'lua', 'php', 'julia', 'Rscript'}


def splitenv(value: str, rest: list[str]) -> list[str]:
    packed = ' '.join([value, *rest]).strip()

    return shlex.split(packed)


def args(args: list[str], *, allow_split: bool=True) -> list[str]:
    i = 0

    while i < len(args):
        arg = args[i]

        if arg == '--':
            return args[i + 1:]

        if allow_split:
            if arg == '-S':
                if i + 1 >= len(args):
                    return []

                return args(splitenv(' '.join(args[i + 1:]), []), allow_split=False)

            if arg.startswith('-S') and len(arg) > 2:
                return args(splitenv(arg[2:], args[i + 1:]), allow_split=False)

            if arg == '-vS':
                if i + 1 >= len(args):
                    return []

                return args(splitenv(' '.join(args[i + 1:]), []), allow_split=False)

            if arg.startswith('-vS') and len(arg) > 3:
                return args(splitenv(arg[3:], args[i + 1:]), allow_split=False)

            if arg.startswith('--split-string='):
                return args(splitenv(arg.split('=', 1)[1], args[i + 1:]), allow_split=False)

            if arg == '--split-string':
                if i + 1 >= len(args):
                    return []

                return args(splitenv(args[i + 1], args[i + 2:]), allow_split=False)

        if arg in {'-u', '-C', '-P', '-a', '--unset', '--chdir', '--argv0'}:
            if i + 2 > len(args):
                return []

            i += 2
            continue

        if arg.startswith(('-u', '-C', '-P', '-a')) and len(arg) > 2 and (not arg.startswith('--')):
            i += 1
            continue

        if arg.startswith(('--unset=', '--chdir=', '--argv0=')):
            i += 1
            continue

        if arg in {'-', '-i', '-0', '-v', '--ignore-environment', '--null', '--debug', '--list-signal-handling'}:
            i += 1
            continue

        if arg.startswith(('--default-signal', '--ignore-signal', '--block-signal')):
            i += 1
            continue

        if arg.startswith('-'):
            return []

        if '=' in arg:
            i += 1
            continue

        return args[i:]

    return []


def shebang(path: Path) -> str | None:
    try:
        with path.open('rb') as f:
            first = f.read(256)

        if not first.startswith(b'#!'):
            return None

        line = first.split(b'\n')[0].decode(errors='replace')[2:].strip()
        parts = shlex.split(line)

        if not parts:
            return None

        interp = Path(parts[0]).name

        if interp == 'env':
            env_args = args(parts[1:])

            if not env_args:
                return None

            interp = Path(env_args[0]).name

        return interp
    except (OSError, ValueError):
        return None


def shebangtype(path: Path) -> FileType | None:
    interp = shebang(path)

    if interp in _SHEBANG_CODE_INTERPRETERS:
        return FileType.CODE

    return None


def classify(path: Path) -> FileType | None:
    from graphify.manifest import match

    if match(path):
        return FileType.CODE

    if path.name.lower().endswith('.blade.php'):
        return FileType.CODE

    ext = path.suffix.lower()

    if not ext:
        return shebangtype(path)

    if ext in CODE_EXTENSIONS:
        return FileType.CODE

    if ext in PAPER_EXTENSIONS:
        if any((part.endswith(tuple(_ASSET_DIR_MARKERS)) for part in path.parts)):
            return None

        return FileType.PAPER

    if ext in IMAGE_EXTENSIONS:
        return FileType.IMAGE

    if ext in DOC_EXTENSIONS:
        if paper(path):
            return FileType.PAPER

        return FileType.DOCUMENT

    if ext in OFFICE_EXTENSIONS:
        return FileType.DOCUMENT

    if ext in GOOGLE_WORKSPACE_EXTENSIONS:
        return FileType.DOCUMENT

    if ext in VIDEO_EXTENSIONS:
        return FileType.VIDEO

    return None

_pypdf_missing_warned = False


def pdf(path: Path) -> str:
    if not cap(path):
        return ''

    try:
        from pypdf import PdfReader
    except ImportError:
        global _pypdf_missing_warned

        if not _pypdf_missing_warned:
            _pypdf_missing_warned = True
            print("[graphify] WARNING: PDF text extraction skipped: 'pypdf' is not installed. Install the pdf extra: uv tool install 'graphifyy[pdf]'", file=sys.stderr)

        return ''

    try:
        reader = PdfReader(str(path))
        pages = []

        for page in reader.pages:
            text = page.extract_text()

            if text:
                pages.append(text)

        return '\n'.join(pages)
    except Exception:
        return ''

_W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
_MC_FALLBACK = '{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback'
_DOCX_WRAPPERS = frozenset({f'{_W}sdt', f'{_W}customXml'})
_DOCX_SKIPPED = frozenset({f'{_W}del', f'{_W}moveFrom', f'{_W}rt', f'{_W}txbxContent', _MC_FALLBACK})
_DOCX_RUN_TEXT = tuple((f'{_W}{tag}' for tag in ('br', 'cr', 'noBreakHyphen', 'ptab', 't', 'tab')))


def docxinside(element, stop, tags) -> bool:
    node = element.getparent()

    while node is not None and node is not stop:
        if node.tag in tags:
            return True

        node = node.getparent()

    return False


def docxtext(p) -> str:
    return ''.join((str(element) for element in p.iter(*_DOCX_RUN_TEXT) if not docxinside(element, p, _DOCX_SKIPPED)))


def docxblocks(container):
    for child in container:
        if child.tag == f'{_W}p':
            yield child

            for box in child.iter(f'{_W}txbxContent'):
                if not docxinside(box, child, _DOCX_SKIPPED):
                    yield from docxblocks(box)
        elif child.tag == f'{_W}tbl':
            yield child
        elif child.tag in _DOCX_WRAPPERS:
            content = child.find(f'{_W}sdtContent')

            yield from docxblocks(child if content is None else content)


def docxchildren(parent, tag: str):
    for child in parent:
        if child.tag == f'{_W}{tag}':
            yield child
        elif child.tag in _DOCX_WRAPPERS:
            content = child.find(f'{_W}sdtContent')

            yield from docxchildren(child if content is None else content, tag)


def docxcell(tc) -> str:
    parts = []

    for block in docxblocks(tc):
        if block.tag == f'{_W}p':
            parts.append(docxtext(block))
        else:
            parts.extend((docxcell(cell) for row in docxchildren(block, 'tr') for cell in docxchildren(row, 'tc')))

    return ' '.join((' '.join(part.split()) for part in parts if part.strip())).replace('|', '\\|')


def docx(path: Path) -> str:
    if not zipcap(path):
        return ''

    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        doc = Document(str(path))
        lines = []

        for block in docxblocks(doc.element.body):
            if block.tag == f'{_W}tbl':
                rows = [[docxcell(cell._tc) for cell in row.cells] for row in Table(block, doc).rows]

                if not rows:
                    continue

                header = '| ' + ' | '.join(rows[0]) + ' |'
                sep = '| ' + ' | '.join(('---' for _ in rows[0])) + ' |'
                lines.extend([header, sep])

                for row in rows[1:]:
                    lines.append('| ' + ' | '.join(row) + ' |')
                continue

            para = Paragraph(block, doc)
            style = para.style.name if para.style else ''
            text = docxtext(block).strip()

            if not text:
                lines.append('')
                continue

            if style.startswith('Heading 1'):
                lines.append(f'# {text}')
            elif style.startswith('Heading 2'):
                lines.append(f'## {text}')
            elif style.startswith('Heading 3'):
                lines.append(f'### {text}')
            elif style.startswith('List'):
                lines.append(f'- {text}')
            else:
                lines.append(text)

        return '\n'.join(lines)
    except ImportError:
        return ''
    except Exception:
        return ''


def xlsx(path: Path) -> str:
    if not zipcap(path):
        return ''

    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        sections = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []

            for row in ws.iter_rows(values_only=True):
                if all((cell is None for cell in row)):
                    continue

                rows.append([str(cell) if cell is not None else '' for cell in row])

            if not rows:
                continue

            sections.append(f'## Sheet: {sheet_name}')

            if len(rows) >= 1:
                header = '| ' + ' | '.join(rows[0]) + ' |'
                sep = '| ' + ' | '.join(('---' for _ in rows[0])) + ' |'
                sections.extend([header, sep])

                for row in rows[1:]:
                    sections.append('| ' + ' | '.join(row) + ' |')
        wb.close()

        return '\n'.join(sections)
    except ImportError:
        return ''
    except Exception:
        return ''


def xlsxstructure(path: Path) -> dict:

    def nid(*parts: str) -> str:
        return re.sub('[^a-z0-9_]', '_', '_'.join((p.lower() for p in parts)).strip('_'))

    try:
        import openpyxl
    except ImportError:
        return {'nodes': [], 'edges': []}

    try:
        wb = openpyxl.load_workbook(str(path), read_only=False, data_only=True)
    except Exception:
        return {'nodes': [], 'edges': []}

    stem = re.sub('[^a-z0-9]', '_', path.stem.lower())
    str_path = str(path)
    file_nid = nid(str_path)
    nodes: list[dict] = [{'id': file_nid, 'label': path.name, 'file_type': 'document', 'source_file': str_path, 'source_location': None}]
    edges: list[dict] = []
    seen: set[str] = {file_nid}

    def add(nid: str, label: str) -> None:
        if nid not in seen:
            seen.add(nid)
            nodes.append({'id': nid, 'label': label, 'file_type': 'document', 'source_file': str_path, 'source_location': None})

    def edge(src: str, tgt: str, relation: str) -> None:
        edges.append({'source': src, 'target': tgt, 'relation': relation, 'source_file': str_path, 'source_location': None})

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        sheet_nid = nid(stem, sheet_name)
        add(sheet_nid, f'{sheet_name} (sheet)')
        edge(file_nid, sheet_nid, 'contains')

        if hasattr(ws, 'tables'):
            for tbl in ws.tables.values():
                tbl_nid = nid(stem, sheet_name, tbl.name)
                add(tbl_nid, tbl.name)
                edge(sheet_nid, tbl_nid, 'contains')
                ref = tbl.ref

                if ref:
                    try:
                        from openpyxl.utils import range_boundaries
                        min_col, min_row, max_col, _ = range_boundaries(ref)
                        header_row = list(ws.iter_rows(min_row=min_row, max_row=min_row, min_col=min_col, max_col=max_col, values_only=True))

                        if header_row:
                            for col_name in header_row[0]:
                                if col_name:
                                    col_nid = nid(stem, tbl.name, str(col_name))
                                    add(col_nid, str(col_name))
                                    edge(tbl_nid, col_nid, 'contains')
                    except Exception:
                        pass
        else:
            for row in ws.iter_rows(max_row=1, values_only=True):
                for cell in row:
                    if cell:
                        col_nid = nid(stem, sheet_name, str(cell))
                        add(col_nid, str(cell))
                        edge(sheet_nid, col_nid, 'contains')
                break

    try:
        wb.close()
    except Exception:
        pass

    return {'nodes': nodes, 'edges': edges}


def office(path: Path, out_dir: Path, root: 'Path | None'=None) -> Path | None:
    ext = path.suffix.lower()

    if ext == '.docx':
        text = docx(path)
    elif ext == '.xlsx':
        text = xlsx(path)
    else:
        return None

    if not text.strip():
        return None

    out_dir.mkdir(parents=True, exist_ok=True)
    import hashlib
    import unicodedata

    if root is None:
        root = out_dir.parent.parent

    try:
        key = path.resolve().relative_to(Path(root).resolve()).as_posix()
    except (ValueError, OSError):
        key = str(path.resolve())
    normalized_path = unicodedata.normalize('NFC', key)
    name_hash = hashlib.sha256(normalized_path.encode()).hexdigest()[:8]
    out_path = out_dir / f'{path.stem}_{name_hash}.md'

    try:
        if out_path.exists() and os.stat(ospath(out_path)).st_mtime >= os.stat(ospath(path)).st_mtime:
            return out_path
    except OSError:
        if out_path.exists():
            return out_path

    out_path.write_text(f'<!-- converted from {path.name} -->\n\n{text}', encoding='utf-8')

    return out_path


def words(path: Path) -> int:
    try:
        ext = path.suffix.lower()

        if ext == '.pdf':
            return len(pdf(path).split())

        if ext == '.docx':
            return len(docx(path).split())

        if ext == '.xlsx':
            return len(xlsx(path).split())

        if not stat.S_ISREG(os.stat(ospath(path)).st_mode):
            return 0

        with open(ospath(path), encoding='utf-8', errors='ignore') as f:
            return len(f.read().split())
    except Exception:
        return 0


def regular(path: Path) -> bool:
    try:
        return stat.S_ISREG(os.stat(path).st_mode)
    except OSError:
        return False

_SKIP_DIRS = {'venv', '.venv', 'node_modules', '__pycache__', '.git', 'dist', 'build', 'target', 'site-packages', 'lib64', '.pytest_cache', '.mypy_cache', '.ruff_cache', '.tox', '.nox', '.eggs', '*.egg-info', 'graphify-out', 'lcov-report', 'visual-tests', 'visual-test', '__snapshots__', 'storybook-static', 'dist-protected', '.next', '.nuxt', '.turbo', '.angular', '.idea', '.cache', '.parcel-cache', '.svelte-kit', '.terraform', '.serverless', '.graphify', '.obsidian', '.smart-env', '.worktrees'}
_SKIP_FILES = {'package-lock.json', 'yarn.lock', 'pnpm-lock.yaml', 'Cargo.lock', 'poetry.lock', 'Gemfile.lock', 'composer.lock', 'go.sum', 'go.work.sum', '.graphifyinclude'}
_JS_SNAPSHOT_TEST_ROOTS = frozenset({'__tests__', '__test__'})
_NESTED_SKILL_HOLDERS = frozenset({('.pi', 'agent'), ('.config', 'kilo')})
_GRAPHIFY_INSTALLED_FILES = frozenset({('.agents', 'rules', 'graphify.md'), ('.agents', 'workflows', 'graphify.md'), ('.cursor', 'rules', 'graphify.mdc'), ('.kilo', 'plugins', 'graphify.js'), ('.kiro', 'steering', 'graphify.md'), ('.opencode', 'plugins', 'graphify.js'), ('.windsurf', 'rules', 'graphify.md')})


def installed(path: 'Path') -> bool:
    return path.parts[-3:] in _GRAPHIFY_INSTALLED_FILES

_COVERAGE_ARTIFACT_FILES = frozenset({'lcov.info', 'coverage-final.json', 'coverage-summary.json', 'clover.xml', 'coverage.xml', 'cobertura-coverage.xml', 'jacoco.xml', '.coverage', 'index.html'})
_COVERAGE_ARTIFACT_DIRS = frozenset({'lcov-report', 'html-report'})


def coverage(d: 'Path') -> bool:
    try:
        for name in _COVERAGE_ARTIFACT_FILES:
            if (d / name).is_file():
                return True

        for name in _COVERAGE_ARTIFACT_DIRS:
            if (d / name).is_dir():
                return True
    except OSError:
        pass

    return False

_BUILD_OUTPUT_SUFFIXES = ('.class', '.jar', '.o', '.obj', '.exe', '.dll', '.pyc', '.js.map', '.css.map', '.d.ts')


def markers(d: 'Path') -> bool:
    try:
        if (d / 'production').is_dir():
            return True

        if (d / '_next').is_dir():
            return True

        subdirs: list = []

        with os.scandir(d) as it:
            for entry in it:
                if entry.is_file() and entry.name.lower().endswith(_BUILD_OUTPUT_SUFFIXES):
                    return True

                if entry.is_dir():
                    subdirs.append(entry.path)
        budget = 4000

        for sub in subdirs:
            if budget <= 0:
                break

            with os.scandir(sub) as it:
                for entry in it:
                    budget -= 1

                    if entry.is_file() and entry.name.lower().endswith(_BUILD_OUTPUT_SUFFIXES):
                        return True

                    if budget <= 0:
                        break
    except OSError:
        pass

    return False


def venv(d: 'Path') -> bool:
    try:
        if (d / 'pyvenv.cfg').is_file():
            return True

        if (d / 'bin' / 'activate').is_file() or (d / 'Scripts' / 'activate').is_file():
            return True

        if next(d.glob('lib/python*'), None) is not None:
            return True

        if (d / 'conda-meta').is_dir():
            return True
    except OSError:
        pass

    return False


def noise(part: str, parent: 'Path | None'=None) -> bool:
    if part in _SKIP_DIRS:
        return True

    if part in ('env', '.env') or part.endswith('_env'):
        if parent is None:
            return False

        return venv(parent / part)

    if part == 'coverage':
        if parent is None:
            return False

        return coverage(parent / part)

    if part == 'out':
        if parent is None:
            return False

        return markers(parent / part)

    if part == 'snapshots':
        if parent is None:
            return False

        snap_dir = parent / part

        if parent.name in _JS_SNAPSHOT_TEST_ROOTS:
            return True

        try:
            if next(snap_dir.glob('*.snap'), None) is not None:
                return True
        except OSError:
            pass

        return False

    if part.endswith('_venv'):
        return True

    if part.endswith('.egg-info'):
        return True

    if part == 'worktrees' and parent is not None and parent.name.startswith('.'):
        return True

    if part == 'graphify' and parent is not None:
        if parent.name == '.aider':
            return True

        if parent.name == 'skills':
            holder = parent.parent

            if holder.name.startswith('.') and holder.name != '..':
                return True

            return (holder.parent.name, holder.name) in _NESTED_SKILL_HOLDERS

    return False

_VCS_MARKERS = ('.git', '.hg', '.svn', '_darcs', '.fossil')


def nfc(text: str) -> str:
    return unicodedata.normalize('NFC', text)


def parseline(raw: str) -> str:
    line = raw.rstrip('\n\r')
    line = line.lstrip()

    if not line or line.startswith('#'):
        return ''

    line = re.sub('\\s+#+[^\\\\].*$', '', line)
    line = line.replace('\\#', '#')
    line = re.sub('(?<!\\\\) +$', '', line)

    return nfc(line)


def vcs(start: Path) -> Path | None:
    current = start.resolve()
    home = Path.home()

    while True:
        if any(((current / m).exists() for m in _VCS_MARKERS)):
            return current

        parent = current.parent

        if parent == current or current == home:
            return None

        current = parent


def identity(path: Path) -> str:
    return nfc(os.path.normcase(os.path.abspath(os.fspath(path))))


def tracked(root: Path) -> tuple[set[str], set[str]]:
    root = root.resolve()
    vcs_root = vcs(root)

    if vcs_root is None or not (vcs_root / '.git').exists():
        return (set(), set())

    try:
        proc = subprocess.run(['git', '-C', str(vcs_root), 'ls-files', '-z', '--cached'], stdin=subprocess.DEVNULL, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, check=False, timeout=30, env={**os.environ, 'GIT_OPTIONAL_LOCKS': '0'})
    except (OSError, subprocess.SubprocessError):
        return (set(), set())

    if proc.returncode != 0:
        return (set(), set())

    tracked_files: set[str] = set()
    tracked_dirs: set[str] = set()

    for raw in proc.stdout.split(b'\x00'):
        if not raw:
            continue

        path = Path(os.path.abspath(vcs_root / os.fsdecode(raw)))

        try:
            path.relative_to(root)
        except ValueError:
            continue

        if not regular(path):
            continue

        tracked_files.add(identity(path))
        parent = path.parent

        while parent != root:
            parent_key = identity(parent)

            if parent_key in tracked_dirs:
                break

            tracked_dirs.add(parent_key)
            parent = parent.parent

    return (tracked_files, tracked_dirs)


def exclude(vcs_root: Path) -> Path | None:
    dot_git = vcs_root / '.git'
    git_dir: Path | None = None

    if dot_git.is_dir():
        git_dir = dot_git
    elif dot_git.is_file():
        try:
            content = dot_git.read_text(encoding='utf-8', errors='ignore').strip()
        except OSError:
            content = ''

        if content.startswith('gitdir:'):
            gd = Path(content[len('gitdir:'):].strip())

            if not gd.is_absolute():
                gd = (vcs_root / gd).resolve()
            git_dir = gd
            commondir = gd / 'commondir'

            if commondir.exists():
                try:
                    cd_raw = commondir.read_text(encoding='utf-8', errors='ignore').strip()
                except OSError:
                    cd_raw = ''

                if cd_raw:
                    cd = Path(cd_raw)
                    git_dir = cd if cd.is_absolute() else (gd / cd).resolve()

    if git_dir is None:
        return None

    exclude = git_dir / 'info' / 'exclude'

    return exclude if exclude.is_file() else None

_warned_ignore_encodings: set[str] = set()


def readignore(path: Path) -> str:
    raw = path.read_bytes()

    try:
        return raw.decode('utf-8-sig')
    except UnicodeDecodeError:
        pass

    if raw[:2] in (b'\xff\xfe', b'\xfe\xff'):
        try:
            return raw.decode('utf-16')
        except UnicodeDecodeError:
            pass
    import locale
    import sys as _sys
    fallback = locale.getpreferredencoding(False) or 'latin-1'

    for enc in (fallback, 'latin-1'):
        try:
            text = raw.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue

        key = str(path)

        if key not in _warned_ignore_encodings:
            _warned_ignore_encodings.add(key)
            print(f'[graphify] WARNING: {path} is not valid UTF-8; read it as {enc} instead. Re-save it as UTF-8 — patterns with non-ASCII characters may not match as written.', file=_sys.stderr)

        return text

    return raw.decode('utf-8', errors='ignore')


def dirignore(d: Path, *, gitignore: bool=True) -> list[tuple[Path, str]]:
    patterns: list[tuple[Path, str]] = []

    for fname in ('.gitignore', '.graphifyignore') if gitignore else ('.graphifyignore',):
        ignore_file = d / fname

        if ignore_file.exists():
            for raw in readignore(ignore_file).splitlines():
                line = parseline(raw)

                if line:
                    patterns.append((d, line))

    return patterns


def ignore(root: Path, *, gitignore: bool=True) -> list[tuple[Path, str]]:
    root = root.resolve()
    ceiling = vcs(root) or root
    dirs: list[Path] = []
    current = root

    while True:
        dirs.append(current)

        if current == ceiling:
            break

        current = current.parent
    dirs.reverse()
    patterns: list[tuple[Path, str]] = []
    info_exclude = exclude(ceiling) if gitignore else None

    if info_exclude is not None:
        for raw in readignore(info_exclude).splitlines():
            line = parseline(raw)

            if line:
                patterns.append((ceiling, line))

    for d in dirs:
        patterns.extend(dirignore(d, gitignore=gitignore))

    return patterns

_PARSED_PATTERN_CACHE: dict[str, tuple[bool, bool, bool, str]] = {}
_PARSED_PATTERN_CACHE_MAX = 100000


def parsepattern(pattern: str) -> tuple[bool, bool, bool, str]:
    got = _PARSED_PATTERN_CACHE.get(pattern)

    if got is None:
        negated = pattern.startswith('!')
        raw = pattern[1:] if negated else pattern
        directory_only = raw.endswith('/')
        path_relative = '/' in raw.rstrip('/')
        got = (negated, directory_only, path_relative, raw.strip('/'))

        if len(_PARSED_PATTERN_CACHE) >= _PARSED_PATTERN_CACHE_MAX:
            _PARSED_PATTERN_CACHE.clear()
        _PARSED_PATTERN_CACHE[pattern] = got

    return got

_CASEFOLD_PATHS = os.path.normcase('Aa') != 'Aa'


def lexical(target: Path, target_parts: tuple[str, ...], anchor: Path) -> str | None:
    anchor_parts = anchor.parts

    if not anchor_parts:
        try:
            return nfc(str(target.relative_to(anchor)).replace(os.sep, '/'))
        except ValueError:
            return None

    n = len(anchor_parts)

    if len(target_parts) < n:
        return None

    head = target_parts[:n]

    if head != anchor_parts and (not _CASEFOLD_PATHS or tuple(map(os.path.normcase, head)) != tuple(map(os.path.normcase, anchor_parts))):
        return None

    tail = target_parts[n:]

    if not tail:
        return '.'

    return nfc('/'.join(tail))


def globstar(path_parts: tuple[str, ...], pattern_parts: tuple[str, ...], path_idx: int, pattern_idx: int, memo: dict[tuple[int, int], bool]) -> bool:
    key = (path_idx, pattern_idx)
    cached = memo.get(key)

    if cached is not None:
        return cached

    if pattern_idx == len(pattern_parts):
        result = path_idx == len(path_parts)
    else:
        part = pattern_parts[pattern_idx]

        if part == '**':
            if pattern_idx == len(pattern_parts) - 1:
                result = path_idx < len(path_parts)
            else:
                result = globstar(path_parts, pattern_parts, path_idx, pattern_idx + 1, memo) or (path_idx < len(path_parts) and globstar(path_parts, pattern_parts, path_idx + 1, pattern_idx, memo))
        else:
            result = path_idx < len(path_parts) and fnmatch.fnmatchcase(path_parts[path_idx], part) and globstar(path_parts, pattern_parts, path_idx + 1, pattern_idx + 1, memo)
    memo[key] = result

    return result


def matchpattern(path: str, pattern: str) -> bool:
    path_parts = tuple(path.split('/'))
    pattern_parts = tuple(pattern.split('/'))

    if '**' not in pattern_parts:
        if len(path_parts) != len(pattern_parts):
            return False

        return all((fnmatch.fnmatchcase(pp, qp) for pp, qp in zip(path_parts, pattern_parts)))

    return globstar(path_parts, pattern_parts, 0, 0, {})


def ignored(path: Path, root: Path, patterns: list[tuple[Path, str]], *, _cache: dict[Path, bool] | None=None) -> bool:
    if not patterns:
        return False

    root_nparts = len(root.parts)

    def eval(target: Path) -> bool:
        if _cache is not None and target in _cache:
            return _cache[target]

        target_parts = target.parts
        target_name_nfc: str | None = None
        target_is_dir: bool | None = None
        rel_root_known = False
        rel_root: str | None = None
        rel_by_anchor: dict[Path, tuple[str | None, int]] = {}
        segs_by_rel: dict[str, tuple[list[str], list[str]]] = {}

        def segments(rel: str) -> tuple[list[str], list[str]]:
            got = segs_by_rel.get(rel)

            if got is None:
                parts = rel.split('/')
                prefixes: list[str] = []
                acc = ''

                for part in parts:
                    acc = part if not acc else acc + '/' + part
                    prefixes.append(acc)
                got = (parts, prefixes)
                segs_by_rel[rel] = got

            return got

        def matches(rel: str, p: str, path_relative: bool) -> bool:
            nonlocal target_name_nfc

            if path_relative:
                return matchpattern(rel, p)

            if fnmatch.fnmatch(rel, p):
                return True

            if target_name_nfc is None:
                target_name_nfc = nfc(target.name)

            if fnmatch.fnmatch(target_name_nfc, p):
                return True

            parts, prefixes = segments(rel)

            for part, prefix in zip(parts, prefixes):
                if fnmatch.fnmatch(part, p):
                    return True

                if fnmatch.fnmatch(prefix, p):
                    return True

            return False

        result = False

        for anchor, pattern in patterns:
            negated, directory_only, path_relative, p = parsepattern(pattern)

            if not p:
                continue

            cached_rel = rel_by_anchor.get(anchor)

            if cached_rel is None:
                cached_rel = (lexical(target, target_parts, anchor), len(anchor.parts))
                rel_by_anchor[anchor] = cached_rel
            rel_anchor, anchor_nparts = cached_rel

            if rel_anchor is None:
                continue

            matched = False

            if rel_anchor != '.':
                rel = rel_anchor

                if not path_relative and root_nparts > anchor_nparts:
                    if not rel_root_known:
                        rel_root_known = True
                        rel_root = lexical(target, target_parts, root)

                    if rel_root is not None:
                        rel = rel_root
                matched = matches(rel, p, path_relative=path_relative)

                if matched and directory_only:
                    if target_is_dir is None:
                        target_is_dir = target.is_dir()

                    if not target_is_dir:
                        matched = False

            if matched:
                result = not negated

        if _cache is not None:
            _cache[target] = result

        return result

    try:
        rel_parts = path.relative_to(root).parts
    except ValueError:
        return eval(path)

    ancestor = root

    for part in rel_parts[:-1]:
        ancestor = ancestor / part

        if eval(ancestor):
            return True

    return eval(path)


def scanignored(path: Path, root: Path, patterns: list[tuple[Path, str]], explicit_patterns: list[tuple[Path, str]], tracked_files: set[str], tracked_dirs: set[str], *, cache: dict[Path, bool], explicit_cache: dict[Path, bool]) -> bool:
    if not ignored(path, root, patterns, _cache=cache):
        return False

    if ignored(path, root, explicit_patterns, _cache=explicit_cache):
        return True

    identity = identity(path)

    return identity not in tracked_files and identity not in tracked_dirs


def predicate(root: Path, *, extra_excludes: list[str] | None=None, gitignore: bool=True) -> Callable[[Path], bool]:
    root = root.resolve()
    patterns = ignore(root, gitignore=gitignore)
    explicit_patterns = ignore(root, gitignore=False)
    tracked_files, tracked_dirs = tracked(root) if gitignore and len(patterns) > len(explicit_patterns) else (set(), set())

    if extra_excludes:
        for pat in extra_excludes:
            line = parseline(pat)

            if line:
                patterns.append((root, line))
                explicit_patterns.append((root, line))
    cache: dict[Path, bool] = {}
    explicit_cache: dict[Path, bool] = {}
    loaded_dirs: set[Path] = {root}

    def ignored(path: Path) -> bool:
        path = Path(os.path.abspath(path))

        try:
            rel_parts = path.relative_to(root).parts
        except ValueError:
            return False

        if path.name in _SKIP_FILES or installed(path):
            return True

        parent = root

        for part in rel_parts[:-1]:
            if noise(part, parent):
                return True

            parent = parent / part
        ancestor = root

        for part in rel_parts[:-1]:
            ancestor = ancestor / part

            if ancestor not in loaded_dirs:
                loaded_dirs.add(ancestor)
                patterns.extend(dirignore(ancestor, gitignore=gitignore))
                explicit_patterns.extend(dirignore(ancestor, gitignore=False))

        return scanignored(path, root, patterns, explicit_patterns, tracked_files, tracked_dirs, cache=cache, explicit_cache=explicit_cache)

    return _ignored


def symlinks(root: Path) -> bool:
    try:
        for p in root.iterdir():
            if p.is_symlink():
                return True
    except (OSError, PermissionError):
        pass

    return False


def resolves(path: Path, root: Path) -> bool:
    try:
        path.resolve().relative_to(root.resolve())
    except (OSError, RuntimeError, ValueError):
        return False

    return True


def detect(root: Path, *, follow_symlinks: bool | None=None, google_workspace: bool | None=None, extra_excludes: list[str] | None=None, cache_root: Path | None=None, gitignore: bool=True) -> dict:
    root = root.resolve()
    configured_out_dir = root / GRAPHIFY_OUT
    configured_out_names = {configured_out_dir.name}

    try:
        configured_out_dir = configured_out_dir.resolve()
    except (OSError, RuntimeError):
        configured_out_dir = configured_out_dir.absolute()
    configured_out_names.add(configured_out_dir.name)

    if (root / '.graphifyinclude').is_file():
        import sys as _sys
        print('[graphify] WARNING: .graphifyinclude is no longer supported (it has been non-functional since dot directories became indexed by default); to re-include ignored paths, use ! negation patterns in .graphifyignore.', file=_sys.stderr)

    if follow_symlinks is None:
        follow_symlinks = False
    google_workspace = enabled() if google_workspace is None else google_workspace
    files: dict[FileType, list[str]] = {FileType.CODE: [], FileType.DOCUMENT: [], FileType.PAPER: [], FileType.IMAGE: [], FileType.VIDEO: []}
    total_words = 0

    def wc(path: Path) -> int:
        from graphify import cache as _cache

        return _cache.count(path, root, words, cache_root=cache_root)

    skipped_sensitive: list[str] = []
    unclassified: list[str] = []
    ignored: list[str] = []
    pruned_noise: list[str] = []
    ignore_patterns = ignore(root, gitignore=gitignore)
    explicit_ignore_patterns = ignore(root, gitignore=False)
    tracked_files, tracked_dirs = tracked(root) if gitignore and len(ignore_patterns) > len(explicit_ignore_patterns) else (set(), set())
    ignore_cache: dict[Path, bool] = {}
    explicit_ignore_cache: dict[Path, bool] = {}

    if extra_excludes:
        for pat in extra_excludes:
            line = parseline(pat)

            if line:
                ignore_patterns.append((root, line))
                explicit_ignore_patterns.append((root, line))

    def scanignored(path: Path) -> bool:
        return scanignored(path, root, ignore_patterns, explicit_ignore_patterns, tracked_files, tracked_dirs, cache=ignore_cache, explicit_cache=explicit_ignore_cache)

    memory_dir = root / GRAPHIFY_OUT / 'memory'
    scan_paths = [root]

    if memory_dir.exists():
        scan_paths.append(memory_dir)
    seen: set[Path] = set()
    all_files: list[Path] = []
    walk_errors: list[str] = []

    def walkerror(err: OSError) -> None:
        import sys as _sys
        target = getattr(err, 'filename', None) or '<unknown>'
        walk_errors.append(f'{target}: {err}')
        print(f"[graphify] WARNING: could not scan {target} ({err}); its files are missing from this run's enumeration.", file=_sys.stderr)

    for scan_root in scan_paths:
        in_memory_tree = memory_dir.exists() and str(scan_root).startswith(str(memory_dir))

        for dirpath, dirnames, filenames in os.walk(scan_root, followlinks=follow_symlinks, onerror=walkerror):
            dp = Path(dirpath)

            if follow_symlinks and os.path.islink(dirpath):
                real = os.path.realpath(dirpath)
                parent_real = os.path.realpath(os.path.dirname(dirpath))

                if parent_real == real or parent_real.startswith(real + os.sep):
                    dirnames.clear()
                    continue

            if not in_memory_tree:
                if dp != root:
                    ignore_patterns.extend(dirignore(dp, gitignore=gitignore))
                    explicit_ignore_patterns.extend(dirignore(dp, gitignore=False))
                kept_dirs: list[str] = []

                for d in dirnames:
                    child = dp / d
                    is_configured_out = False

                    if d in configured_out_names:
                        try:
                            is_configured_out = child.resolve() == configured_out_dir
                        except (OSError, RuntimeError):
                            pass

                    if is_configured_out:
                        pruned_noise.append(str(child) + os.sep)
                        continue

                    if noise(d, dp):
                        pruned_noise.append(str(child) + os.sep)
                        continue

                    if scanignored(child):
                        ignored.append(str(child) + os.sep)
                        continue

                    kept_dirs.append(d)
                dirnames[:] = kept_dirs

                if follow_symlinks:
                    safe_dirs: list[str] = []

                    for d in dirnames:
                        child = dp / d

                        if child.is_symlink() and (not resolves(child, root)):
                            skipped_sensitive.append(str(child) + ' [symlink target outside scan root]')
                            continue

                        safe_dirs.append(d)
                    dirnames[:] = safe_dirs

            for fname in filenames:
                if fname in _SKIP_FILES:
                    continue

                p = dp / fname

                if installed(p):
                    continue

                if p not in seen:
                    seen.add(p)
                    all_files.append(p)
    all_files.sort(key=lambda p: str(p))
    out_base = Path(cache_root).resolve() if cache_root is not None else root
    converted_dir = out_base / GRAPHIFY_OUT / 'converted'

    for p in all_files:
        in_memory = memory_dir.exists() and str(p).startswith(str(memory_dir))

        if not in_memory:
            if str(p).startswith(str(converted_dir)):
                continue

        if not in_memory and scanignored(p):
            ignored.append(str(p))
            continue

        if not resolves(p, root):
            skipped_sensitive.append(str(p) + ' [symlink target outside scan root]')
            continue

        if not regular(p):
            skipped_sensitive.append(str(p) + ' [not a regular file]')
            continue

        if sensitive(p):
            skipped_sensitive.append(str(p))
            continue

        ftype = classify(p)

        if not ftype:
            unclassified.append(str(p))
            continue

        if ftype:
            if p.suffix.lower() in GOOGLE_WORKSPACE_EXTENSIONS:
                if not google_workspace:
                    skipped_sensitive.append(str(p) + ' [Google Workspace shortcut skipped - pass --google-workspace or set GRAPHIFY_GOOGLE_WORKSPACE=1]')
                    continue

                try:
                    md_path = workspace(p, converted_dir, xlsx=xlsx, root=root)
                except Exception as exc:
                    skipped_sensitive.append(str(p) + f' [Google Workspace export failed: {exc}]')
                    continue

                if md_path:
                    if scanignored(md_path) and (not md_path.is_relative_to(converted_dir)):
                        continue

                    files[ftype].append(str(md_path))
                    total_words += wc(md_path)
                else:
                    skipped_sensitive.append(str(p) + ' [Google Workspace export produced no readable text]')
                continue

            if p.suffix.lower() in OFFICE_EXTENSIONS:
                md_path = office(p, converted_dir, root=root)

                if md_path:
                    if scanignored(md_path) and (not md_path.is_relative_to(converted_dir)):
                        continue

                    files[ftype].append(str(md_path))
                    total_words += wc(md_path)
                else:
                    skipped_sensitive.append(str(p) + ' [office conversion failed - pip install graphifyy[office]]')
                continue

            files[ftype].append(str(p))

            if ftype != FileType.VIDEO:
                total_words += wc(p)

    for ftype in files:
        files[ftype].sort()
    total_files = sum((len(v) for v in files.values()))
    needs_graph = total_words >= CORPUS_WARN_THRESHOLD
    warning: str | None = None

    if not needs_graph:
        warning = f'Corpus is ~{total_words:,} words - fits in a single context window. You may not need a graph.'
    elif total_words >= CORPUS_UPPER_THRESHOLD or total_files >= FILE_COUNT_UPPER:
        warning = f'Large corpus: {total_files} files · ~{total_words:,} words. Consider running on a subfolder.'

    return {'files': {k.value: v for k, v in files.items()}, 'total_files': total_files, 'total_words': total_words, 'needs_graph': needs_graph, 'warning': warning, 'skipped_sensitive': skipped_sensitive, 'unclassified': sorted(unclassified), 'walk_errors': walk_errors, 'ignored': sorted(ignored), 'pruned_noise_dirs': sorted(pruned_noise), 'graphifyignore_patterns': len(ignore_patterns), 'scan_root': str(root.resolve())}


def ospath(path: Path) -> str:
    import sys

    if sys.platform != 'win32':
        return str(path)

    s = str(path)

    if s.startswith('\\\\?\\'):
        return s

    try:
        s = os.path.abspath(s)
    except Exception:
        return str(path)

    if s.startswith('\\\\'):
        return '\\\\?\\UNC\\' + s[2:]

    return '\\\\?\\' + s


def md5(path: Path) -> str:
    import hashlib as _hl
    h = _hl.md5(usedforsecurity=False)

    try:
        with open(ospath(path), 'rb') as f:
            for chunk in iter(lambda: f.read(65536), b''):
                h.update(chunk)
    except OSError:
        return ''

    return h.hexdigest()


def stathash(path_str: str) -> tuple[str, float, str] | None:
    try:
        p = Path(path_str)

        return (path_str, os.stat(ospath(p)).st_mtime, md5(p))
    except OSError:
        return None


def nfc(s: str) -> str:
    import unicodedata

    return unicodedata.normalize('NFC', s)


def relstorage(key: str, root: Path) -> str:
    p = Path(key)

    if not p.is_absolute():
        return key

    try:
        base = nfc(str(Path(root).resolve()))
        rel = os.path.relpath(nfc(str(p)), base)
    except (ValueError, OSError):
        return key

    if rel == '..' or rel.startswith('..' + os.sep) or rel.startswith('../'):
        return key

    return rel.replace(os.sep, '/')

_DRIVE_LETTER_RE = re.compile('^[A-Za-z]:[\\\\/]')


def isabs(key: str) -> bool:
    return Path(key).is_absolute() or key.startswith(('/', '\\')) or bool(_DRIVE_LETTER_RE.match(key))


def normflavor(key: str) -> str:
    if Path(key).is_absolute():
        return os.path.normpath(key)

    if key.startswith('\\') or _DRIVE_LETTER_RE.match(key):
        return ntpath.normpath(key)

    return posixpath.normpath(key)


def absstorage(key: str, root: Path) -> str:
    if isabs(key):
        return normflavor(key)

    return nfc(os.path.normpath(str(Path(root).resolve() / Path(key))))


def collapse(items: Iterable[tuple[str, Any]], key_fn: Callable[[str], str]) -> dict[str, Any]:
    result: dict[str, Any] = {}

    for k, v in items:
        canonical = key_fn(k)

        if canonical in result:
            prev = result[canonical]
            prev_seen = prev.get('seen') if isinstance(prev, dict) else None
            new_seen = v.get('seen') if isinstance(v, dict) else None

            if isinstance(prev_seen, (int, float)) and isinstance(new_seen, (int, float)) and (new_seen < prev_seen):
                continue

        result[canonical] = v

    return result


def anchor(manifest_path: str | Path | None, root: Path | None) -> Path | None:
    if root is None or not manifest_path:
        return root

    try:
        mp = Path(manifest_path).resolve()

        if mp.parent.name in ('graphify-out', GRAPHIFY_OUT_NAME):
            manifest_base = mp.parent.parent
            root_resolved = Path(root).resolve()

            if root_resolved != manifest_base:
                try:
                    root_resolved.relative_to(manifest_base)

                    return manifest_base
                except ValueError:
                    pass
    except (OSError, RuntimeError):
        pass

    return root


def load(manifest_path: str=_MANIFEST_PATH, *, root: Path | None=None) -> dict:
    try:
        raw = json.loads(Path(manifest_path).read_text(encoding='utf-8'))
    except Exception:
        return {}

    if not isinstance(raw, dict):
        return raw

    if root is None:
        return collapse(raw.items(), nfc)

    anchor = anchor(manifest_path, root) or root

    return collapse(raw.items(), lambda k: nfc(absstorage(k, anchor)))


def prefix(abs_key: str, relative_keys: set[str]) -> str | None:
    parts = nfc(abs_key).replace('\\', '/').split('/')
    best: tuple[int, str] | None = None

    for i in range(1, len(parts)):
        suffix = '/'.join(parts[i:])

        if suffix not in relative_keys:
            continue

        prefix = '/'.join(parts[:i])

        if prefix in ('', '/', '\\'):
            continue

        if not Path(prefix).is_dir():
            continue

        if best is None or len(suffix) > best[0]:
            best = (len(suffix), prefix)

    return None if best is None else best[1]


def keys(existing_keys: Iterable[Any], storage_root: Path | None, relative_keys: set[str]) -> set[str]:
    if storage_root is None:
        return set()

    groups: dict[str, set[str]] = {}

    for key in existing_keys:
        if not isinstance(key, str) or not isabs(key):
            continue

        stored = nfc(relstorage(key, storage_root))

        if not isabs(stored):
            continue

        prefix = prefix(stored, relative_keys)

        if prefix is None:
            continue

        groups.setdefault(prefix, set()).add(key)
    stale: set[str] = set()

    for keys in groups.values():
        if len(keys) >= 2:
            stale.update(keys)

    return stale


def save(files: dict[str, list[str]], manifest_path: str=_MANIFEST_PATH, *, kind: str='both', root: Path | None=None, scan_corpus: set[str] | list[str] | None=None, clear_semantic: set[str] | list[str] | None=None, clear_ast: set[str] | list[str] | None=None) -> None:
    existing = load(manifest_path, root=root)

    def pathindex(paths: set[str] | list[str] | None) -> set[str] | None:
        if paths is None:
            return None

        indexed: set[str] = set()

        for p in paths:
            indexed.add(p)
            indexed.add(nfc(p))

        return indexed

    scan_set = pathindex(scan_corpus)
    clear_set = pathindex(clear_semantic)
    clear_ast_set = pathindex(clear_ast)

    try:
        root_res: Path | None = Path(root).resolve() if root is not None else None
    except (OSError, RuntimeError):
        root_res = Path(root) if root is not None else None

    def inscan(path_str: str) -> bool:
        if path_str in scan_set or nfc(path_str) in scan_set:
            return True

        try:
            resolved = str(Path(path_str).resolve())

            return resolved in scan_set or nfc(resolved) in scan_set
        except (OSError, RuntimeError):
            return False

    def inclear(path_str: str) -> bool:
        if clear_set is None:
            return False

        if path_str in clear_set or nfc(path_str) in clear_set:
            return True

        try:
            resolved = str(Path(path_str).resolve())

            return resolved in clear_set or nfc(resolved) in clear_set
        except (OSError, RuntimeError):
            return False

    def inclearast(path_str: str) -> bool:
        if clear_ast_set is None:
            return False

        if path_str in clear_ast_set or nfc(path_str) in clear_ast_set:
            return True

        try:
            resolved = str(Path(path_str).resolve())

            return resolved in clear_ast_set or nfc(resolved) in clear_ast_set
        except (OSError, RuntimeError):
            return False

    def inroot(path_str: str) -> bool:
        if root_res is None:
            return False

        p = Path(path_str)

        try:
            p.relative_to(root_res)

            return True
        except ValueError:
            pass

        try:
            p.resolve().relative_to(root_res)

            return True
        except (ValueError, OSError, RuntimeError):
            return False

    def normentry(entry):
        if isinstance(entry, (int, float)):
            return {'mtime': entry, 'ast_hash': '', 'semantic_hash': ''}

        if isinstance(entry, dict) and 'hash' in entry and ('ast_hash' not in entry):
            return {'mtime': entry.get('mtime', 0), 'ast_hash': entry['hash'], 'semantic_hash': ''}

        if isinstance(entry, dict):
            return entry

        return None

    all_files = [f for file_list in files.values() for f in file_list]
    storage_root = anchor(manifest_path, root) if root is not None else None
    stale_checkout: set[str] = set()

    if scan_set is not None and storage_root is not None:
        new_rels: set[str] = set()

        for f in all_files:
            stored = nfc(relstorage(nfc(f), storage_root))

            if not isabs(stored):
                new_rels.add(stored)
        stale_checkout = keys(existing.keys(), storage_root, new_rels)
    manifest: dict[str, dict] = {}

    for f, entry in existing.items():
        normalised = normentry(entry)

        if normalised is None:
            continue

        try:
            if not Path(f).exists():
                continue
        except OSError:
            continue

        if f in stale_checkout:
            continue

        if scan_set is not None and (not inscan(f)) and inroot(f):
            continue

        if clear_ast_set is not None and inclearast(f):
            normalised = {**normalised, 'ast_hash': '', 'semantic_hash': ''}
        elif clear_set is not None and inclear(f):
            normalised = {**normalised, 'semantic_hash': ''}
        manifest[f] = normalised

    with ThreadPoolExecutor() as pool:
        raw = pool.map(stathash, all_files)
    hashed: dict[str, tuple[float, str]] = {r[0]: (r[1], r[2]) for r in raw if r is not None}

    for f in all_files:
        if f not in hashed:
            continue

        mtime, h = hashed[f]
        key = nfc(f)
        prev = normentry(existing.get(key, {})) or {}

        if kind in ('ast', 'both'):
            ast_h = h
        else:
            ast_h = prev.get('ast_hash', '')

        if kind in ('semantic', 'both'):
            sem_h = h
        else:
            sem_h = prev.get('semantic_hash', '') if h == prev.get('ast_hash', '') else ''
        prev_seen = prev.get('seen')
        is_unchanged = isinstance(prev_seen, (int, float)) and mtime == prev.get('mtime') and (ast_h == prev.get('ast_hash', '') if kind in ('ast', 'both') else True) and (sem_h == prev.get('semantic_hash', '') if kind in ('semantic', 'both') else True) and (not inclearast(f)) and (not inclear(f))
        entry: dict = {'mtime': mtime, 'seen': prev_seen if is_unchanged else time.time(), 'ast_hash': ast_h, 'semantic_hash': sem_h}
        manifest[key] = entry

    if storage_root is not None:
        manifest = collapse(manifest.items(), lambda k: nfc(relstorage(k, storage_root)))
    else:
        manifest = collapse(manifest.items(), nfc)
    manifest_p = Path(manifest_path)

    if manifest_p.is_file():
        try:
            disk_raw = json.loads(manifest_p.read_text(encoding='utf-8'))

            if isinstance(disk_raw, dict) and disk_raw == manifest:
                return
        except Exception:
            pass
    from graphify.paths import json as write_json
    write_json(manifest_path, manifest, indent=2)


def hiderewrite(current_mtime: float, stored: dict) -> bool:
    seen = stored.get('seen')

    if not isinstance(seen, (int, float)):
        return False

    delta = float(seen) - float(current_mtime)

    if delta < 0:
        return False

    coarse = float(current_mtime).is_integer()

    return delta < (_MTIME_COARSE_S if coarse else _MTIME_SUBSECOND_S)


def incremental(root: Path, manifest_path: str=_MANIFEST_PATH, *, follow_symlinks: bool | None=None, google_workspace: bool | None=None, kind: str='semantic', extra_excludes: list[str] | None=None, gitignore: bool=True, cache_root: Path | None=None) -> dict:
    full = detect(root, follow_symlinks=follow_symlinks, google_workspace=google_workspace, extra_excludes=extra_excludes, gitignore=gitignore, cache_root=cache_root)
    manifest = load(manifest_path, root=root)

    if not manifest:
        full['incremental'] = True
        full['new_files'] = full['files']
        full['unchanged_files'] = {k: [] for k in full['files']}
        full['new_total'] = full['total_files']
        full['deleted_files'] = []
        full['excluded_files'] = []

        return full

    new_files: dict[str, list[str]] = {k: [] for k in full['files']}
    unchanged_files: dict[str, list[str]] = {k: [] for k in full['files']}

    for ftype, file_list in full['files'].items():
        for f in file_list:
            stored = manifest.get(nfc(f))

            try:
                current_mtime = os.stat(ospath(Path(f))).st_mtime
            except Exception:
                current_mtime = 0

            if isinstance(stored, (int, float)):
                changed = current_mtime != stored
            elif isinstance(stored, dict):
                if 'hash' in stored and 'ast_hash' not in stored:
                    stored = {'mtime': stored.get('mtime', 0), 'ast_hash': stored['hash'], 'semantic_hash': ''}
                hash_key = 'semantic_hash' if kind == 'semantic' else 'ast_hash'
                stored_hash = stored.get(hash_key, '')

                if not stored_hash:
                    changed = True
                else:
                    stored_mtime = stored.get('mtime')

                    if isinstance(stored_mtime, dict):
                        stored_mtime = stored_mtime.get('mtime')

                    if not isinstance(stored_mtime, (int, float)):
                        stored_mtime = None

                    if stored_mtime is None or current_mtime != stored_mtime:
                        changed = md5(Path(f)) != stored_hash
                    elif hiderewrite(current_mtime, stored):
                        changed = md5(Path(f)) != stored_hash
                    else:
                        changed = False
            else:
                changed = True

            if changed:
                new_files[ftype].append(f)
            else:
                unchanged_files[ftype].append(f)

    try:
        root_res: Path | None = Path(root).resolve() if root is not None else None
    except (OSError, RuntimeError):
        root_res = Path(root) if root is not None else None

    def inroot(path_str: str) -> bool:
        if root_res is None:
            return True

        p = Path(path_str)

        try:
            p.relative_to(root_res)

            return True
        except ValueError:
            pass

        try:
            p.resolve().relative_to(root_res)

            return True
        except (ValueError, OSError, RuntimeError):
            return False

    current_files = {nfc(f) for flist in full['files'].values() for f in flist}
    deleted_files: list[str] = []
    excluded_files: list[str] = []

    for f in manifest:
        if not inroot(f):
            continue

        if nfc(f) in current_files:
            continue

        try:
            alive = Path(f).exists()
        except OSError:
            alive = False
        (excluded_files if alive else deleted_files).append(f)
    new_total = sum((len(v) for v in new_files.values()))
    full['incremental'] = True
    full['new_files'] = new_files
    full['unchanged_files'] = unchanged_files
    full['new_total'] = new_total
    full['deleted_files'] = deleted_files
    full['excluded_files'] = excluded_files

    return full